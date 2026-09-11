"""Konverter Markdown -> Word (.docx) untuk manual book.

Dipakai untuk mengubah docs/PANDUAN-PENGGUNA.md menjadi berkas .docx yang rapi
(judul, heading bertingkat, tabel, bullet, teks tebal). Sengaja ditulis
sederhana dan hanya menangani subset Markdown yang dipakai di manual kita:

- Heading: '#', '##', '###'
- Tabel pipa: '| a | b |' dengan baris pemisah '|---|---|'
- Bullet list: baris diawali '- '
- Ordered list: baris diawali '1. ', '2. ', dst.
- Bold inline: '**teks**'
- Blockquote: baris diawali '> '
- Horizontal rule: '---' (dilewati / jadi pemisah)
- Inline code: `teks` (ditampilkan apa adanya tanpa backtick)

Jalankan:
    .venv/bin/python scripts/md_to_docx.py docs/PANDUAN-PENGGUNA.md docs/PANDUAN-PENGGUNA.docx
"""

from __future__ import annotations

import re
import sys
from typing import List

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_CODE_RE = re.compile(r"`([^`]+)`")


def _add_runs_with_inline(paragraph, text: str) -> None:
    """Tambahkan text ke paragraph sambil menangani **bold** dan `code`.

    Backtick inline dihilangkan (isinya ditampilkan sebagai teks biasa).
    """
    # Hilangkan backtick inline lebih dulu (tanpa mengubah isi).
    text = _CODE_RE.sub(lambda m: m.group(1), text)

    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _is_table_separator(line: str) -> bool:
    """True bila baris adalah pemisah header tabel, mis. '|---|:--:|'."""
    stripped = line.strip()
    if "|" not in stripped:
        return False
    cells = [c.strip() for c in stripped.strip("|").split("|")]
    return all(set(c) <= set("-: ") and c != "" for c in cells)


def _split_row(line: str) -> List[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path: str, docx_path: str) -> None:
    with open(md_path, "r", encoding="utf-8") as fh:
        lines = fh.readlines()

    doc = Document()

    # Sedikit penyesuaian gaya default agar enak dibaca.
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i].rstrip("\n")
        line = raw.strip()

        # Baris kosong -> lewati (spacing diatur oleh style paragraf).
        if not line:
            i += 1
            continue

        # Horizontal rule -> lewati.
        if line == "---":
            i += 1
            continue

        # Heading.
        if line.startswith("### "):
            doc.add_heading(_strip_inline(line[4:]), level=3)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(_strip_inline(line[3:]), level=2)
            i += 1
            continue
        if line.startswith("# "):
            # Judul dokumen (level 0 = Title).
            title = doc.add_heading(_strip_inline(line[2:]), level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        # Tabel: baris saat ini punya '|' dan baris berikutnya adalah separator.
        if "|" in line and i + 1 < n and _is_table_separator(lines[i + 1]):
            header = _split_row(line)
            rows: List[List[str]] = []
            j = i + 2
            while j < n and "|" in lines[j] and lines[j].strip():
                rows.append(_split_row(lines[j].strip()))
                j += 1

            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            hdr_cells = table.rows[0].cells
            for idx, htext in enumerate(header):
                hdr_cells[idx].text = ""
                p = hdr_cells[idx].paragraphs[0]
                run = p.add_run(_strip_inline(htext))
                run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for idx in range(len(header)):
                    val = row[idx] if idx < len(row) else ""
                    cells[idx].text = ""
                    _add_runs_with_inline(cells[idx].paragraphs[0], val)
            doc.add_paragraph()  # jeda setelah tabel
            i = j
            continue

        # Blockquote.
        if line.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs_with_inline(p, line[2:])
            i += 1
            continue

        # Ordered list (1. 2. 3. ...).
        m_ol = re.match(r"^\d+\.\s+(.*)$", line)
        if m_ol:
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_inline(p, m_ol.group(1))
            i += 1
            continue

        # Bullet list.
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_inline(p, line[2:])
            i += 1
            continue

        # Paragraf biasa.
        p = doc.add_paragraph()
        _add_runs_with_inline(p, line)
        i += 1

    doc.save(docx_path)
    print("Tersimpan:", docx_path)


def _strip_inline(text: str) -> str:
    """Bersihkan penanda inline untuk teks yang tidak butuh run terpisah."""
    text = _CODE_RE.sub(lambda m: m.group(1), text)
    text = _BOLD_RE.sub(lambda m: m.group(1), text)
    return text


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python scripts/md_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
